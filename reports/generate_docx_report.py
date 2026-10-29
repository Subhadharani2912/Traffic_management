import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()

    # Define color palette
    NAVY = RGBColor(27, 54, 93)       # #1B365D - Kongu Blue / Professional Academic
    DARK_GRAY = RGBColor(44, 62, 80)  # #2C3E50
    BODY_COLOR = RGBColor(30, 30, 30) # #1E1E1E
    WHITE = RGBColor(255, 255, 255)
    
    # -------------------------------------------------------------
    # Helper Functions for XML Manipulation & Styling
    # -------------------------------------------------------------
    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color=BODY_COLOR):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color

    def add_p(container, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5, bold_prefix=""):
        p = container.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            format_run(r_bold, bold=True)
        if text:
            r_text = p.add_run(text)
            format_run(r_text)
        return p

    def add_chapter_title(doc, chapter_num, title):
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(24)
        p1.paragraph_format.space_after = Pt(6)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(f"CHAPTER {chapter_num}\n")
        format_run(r1, size_pt=16, bold=True, color=NAVY)

        r2 = p1.add_run(title.upper())
        format_run(r2, size_pt=16, bold=True, color=NAVY)
        return p1

    def add_heading_2(doc, title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        format_run(r, size_pt=13.5, bold=True, color=NAVY)
        return p

    def add_heading_3(doc, title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        format_run(r, size_pt=12, bold=True, color=DARK_GRAY)
        return p

    def add_heading_4(doc, title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        format_run(r, size_pt=12, bold=True, italic=True, color=DARK_GRAY)
        return p

    def add_fig(doc, img_rel_path, fig_num, title, width_in=5.8):
        img_full = os.path.join(r"c:\Users\ssubh\OneDrive\mini\Traffic_management", img_rel_path)
        if os.path.exists(img_full):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(10)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_full, width=Inches(width_in))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(14)
            r = p_cap.add_run(f"Figure {fig_num}: {title}")
            format_run(r, size_pt=10.5, bold=True, italic=True, color=DARK_GRAY)
        else:
            print(f"Warning: Image not found at {img_full}")

    def add_styled_table(doc, table_num, title, headers, data, col_widths=None):
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(12)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(f"Table {table_num}: {title}")
        format_run(r_cap, size_pt=11, bold=True, color=DARK_GRAY)

        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="B0C4DE", sz="6", val="single")

        # Header Row
        hdr_row = table.rows[0]
        # set header repeat across pages
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for i, header_text in enumerate(headers):
            cell = hdr_row.cells[i]
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(header_text)
            format_run(r, size_pt=10, bold=True, color=WHITE)

        # Data Rows
        for r_idx, row_data in enumerate(data):
            row = table.rows[r_idx + 1]
            bg_color = "F7F9FC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                cell = row.cells[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(str(cell_value))
                format_run(r, size_pt=10, color=BODY_COLOR)

        if col_widths:
            for row in table.rows:
                for c_idx, width in enumerate(col_widths):
                    row.cells[c_idx].width = Inches(width)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(10)

    # Set Margins on Section 1
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.25)
    sec1.right_margin = Inches(1.0)
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)

    print("Helper functions and page setup complete.")
    return doc

create_report()
